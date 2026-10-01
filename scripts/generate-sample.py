"""Original EquationRescue teaching fixture, generated with python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=Path(__file__).resolve().parents[1]/'public/examples/teaching-handout.docx'
doc=Document();section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7);section.left_margin=section.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2','Caption','Header','Footer']:
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(14)
section.header.paragraphs[0].text='EquationRescue  /  Original demonstration handout'
section.header.paragraphs[0].style='Caption'
doc.add_paragraph('Algebra and patterns practice',style='Title')
doc.add_paragraph('This original handout demonstrates repairing LaTeX source inside an existing Word document. Review each equation, then keep working in the same document layout.')
doc.add_paragraph('Warm up',style='Heading 1')
p=doc.add_paragraph();p.add_run('1. Preserve mixed formatting. ').bold=True;p.add_run('Simplify ');p.add_run(r'\(x').bold=True;p.add_run(r'^2 + 2x + 1\)').italic=True;p.add_run(' and explain the pattern.').italic=True
doc.add_paragraph(r'2. A circle has area \(A = \pi r^2\). What happens when the radius doubles?')
doc.add_paragraph('Reference equations',style='Heading 1')
doc.add_paragraph('Use the quadratic formula to compare your answers.')
doc.add_paragraph(r'\[x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}\]')
doc.add_paragraph(r'$$\sum_{i=1}^{n} i = \frac{n(n+1)}{2}$$')
doc.add_paragraph('Practice in a table',style='Heading 1')
table=doc.add_table(rows=1,cols=2);table.autofit=False;table.columns[0].width=Inches(2.5);table.columns[1].width=Inches(3.8)
for cell,text in zip(table.rows[0].cells,['Prompt','Equation source']):
    cell.text=text
    for run in cell.paragraphs[0].runs:run.bold=True
for a,b in [('A ratio',r'\(\frac{a+b}{c}\)'),('A square root',r'\(y = \sqrt{x^2 + 1}\)'),('An indexed sequence',r'\(a_n = n^2 + \alpha\)')]:
    cells=table.add_row().cells;cells[0].text=a;cells[1].text=b
for row in table.rows:
    for cell in row.cells:
        tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
        for side in ('top','left','bottom','right'):
            edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
        tcPr.append(borders)
        margins=OxmlElement('w:tcMar')
        for side in ('top','left','bottom','right'):
            item=OxmlElement('w:'+side);item.set(qn('w:w'),'90');item.set(qn('w:type'),'dxa');margins.append(item)
        tcPr.append(margins)
        for p in cell.paragraphs:p.paragraph_format.space_after=Pt(3)
doc.add_paragraph('Keep unsupported source visible',style='Heading 1')
doc.add_paragraph('The matrix below is intentionally outside the first release. It should stay as typed, with a clear reason in the review report.')
doc.add_paragraph(r'\[\begin{matrix}1 & 0 \\ 0 & 1\end{matrix}\]')
foot=section.footer.paragraphs[0];foot.text='Original fixture · MIT licensed · Not a full LaTeX or Word compatibility test';foot.style='Caption'
# Remove decorative borders inherited from the bundled template.
for element in [doc._element,doc.styles.element]:
    for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for style in doc.styles:
    if style.type == 1:
        for color in style.element.iter(qn('w:color')):
            for attr in ('themeColor','themeTint','themeShade'):
                color.attrib.pop(qn('w:'+attr),None)
doc.core_properties.title='Algebra and patterns practice';doc.core_properties.author='EquationRescue contributors';doc.core_properties.subject='Original demonstration document for local equation repair';doc.save(out);print(out)
